"""
Product Sheet: the Excel-style product list in the Products tab.
One sheet per kind of shop (Fruits, Vegetables, Kirana, ...), pre-filled with common products,
merged with the owner's own catalog. Also builds and reads the matching .xlsx file.
Pure helpers only; the routes live in inventory_api.py.
"""

import base64
import io
import zipfile

# Sheet name -> common products as (name, unit). The sheet name becomes the product category.
DEFAULT_SHEETS = {
    "Fruits": [
        ("Apple", "kg"), ("Banana", "dozen"), ("Mango", "kg"), ("Orange", "kg"), ("Grapes", "kg"),
        ("Papaya", "piece"), ("Pineapple", "piece"), ("Watermelon", "piece"), ("Pomegranate", "kg"),
        ("Guava", "kg"), ("Chikoo", "kg"), ("Mosambi", "kg"), ("Pear", "kg"), ("Kiwi", "piece"),
        ("Custard Apple", "kg"), ("Strawberry", "box"), ("Coconut", "piece"), ("Dates", "kg"),
    ],
    "Vegetables": [
        ("Potato", "kg"), ("Onion", "kg"), ("Tomato", "kg"), ("Cabbage", "piece"), ("Cauliflower", "piece"),
        ("Brinjal", "kg"), ("Lady Finger", "kg"), ("Green Chilli", "kg"), ("Ginger", "kg"), ("Garlic", "kg"),
        ("Carrot", "kg"), ("Capsicum", "kg"), ("Cucumber", "kg"), ("Bottle Gourd", "piece"),
        ("Spinach", "bunch"), ("Methi", "bunch"), ("Coriander", "bunch"), ("Green Peas", "kg"),
        ("Beans", "kg"), ("Beetroot", "kg"), ("Lemon", "piece"), ("Drumstick", "piece"),
    ],
    "Kirana": [
        ("Rice", "kg"), ("Wheat Atta", "kg"), ("Toor Dal", "kg"), ("Moong Dal", "kg"), ("Chana Dal", "kg"),
        ("Masoor Dal", "kg"), ("Sugar", "kg"), ("Salt", "packet"), ("Jaggery", "kg"), ("Tea Powder", "packet"),
        ("Sunflower Oil", "litre"), ("Groundnut Oil", "litre"), ("Ghee", "litre"), ("Barik Rawa", "kg"),
        ("Jada Rawa", "kg"), ("Poha", "kg"), ("Besan", "kg"), ("Maida", "kg"), ("Turmeric Powder", "packet"),
        ("Red Chilli Powder", "packet"), ("Garam Masala", "packet"), ("Mustard Seeds", "packet"),
        ("Jeera", "packet"), ("Biscuits", "packet"), ("Bath Soap", "piece"), ("Detergent Powder", "packet"),
        ("Toothpaste", "piece"), ("Matchbox", "piece"),
    ],
    "Shoes": [
        ("Men's Formal Shoes", "pair"), ("Men's Sports Shoes", "pair"), ("Men's Sandals", "pair"),
        ("Men's Slippers", "pair"), ("Women's Sandals", "pair"), ("Women's Heels", "pair"),
        ("Women's Flats", "pair"), ("Women's Slippers", "pair"), ("Kids' Shoes", "pair"),
        ("Kids' Sandals", "pair"), ("School Shoes", "pair"), ("Flip-Flops", "pair"), ("Socks", "pair"),
        ("Shoe Polish", "piece"), ("Shoe Laces", "pair"),
    ],
    "Mobile": [
        ("Smartphone", "piece"), ("Feature Phone", "piece"), ("Mobile Charger", "piece"),
        ("Type-C Cable", "piece"), ("Micro USB Cable", "piece"), ("Lightning Cable", "piece"),
        ("Wired Earphones", "piece"), ("Bluetooth Earbuds", "piece"), ("Power Bank", "piece"),
        ("Back Cover", "piece"), ("Tempered Glass", "piece"), ("Memory Card 32GB", "piece"),
        ("Memory Card 64GB", "piece"), ("Pen Drive", "piece"), ("SIM Card", "piece"), ("Mobile Stand", "piece"),
        ("Car Charger", "piece"), ("Bluetooth Speaker", "piece"), ("Smart Watch", "piece"),
    ],
    "Medical": [
        ("Paracetamol 500mg", "strip"), ("Cough Syrup", "bottle"), ("Antacid Tablets", "strip"),
        ("Vitamin C Tablets", "strip"), ("ORS", "packet"), ("Bandage Roll", "piece"), ("Adhesive Bandages", "box"),
        ("Cotton Roll", "piece"), ("Antiseptic Liquid", "bottle"), ("Pain Relief Spray", "piece"),
        ("Pain Relief Balm", "piece"), ("Digital Thermometer", "piece"), ("Hand Sanitizer", "bottle"),
        ("Face Mask", "box"), ("Glucose Powder", "packet"), ("Eye Drops", "bottle"), ("Multivitamin Tablets", "strip"),
    ],
}

# Which sheet to open first for each kind of business (matched against the business type name).
BUSINESS_SHEET_HINTS = [
    ("fruit", "Fruits"), ("vegetable", "Vegetables"), ("kirana", "Kirana"), ("retail", "Kirana"),
    ("shoe", "Shoes"), ("footwear", "Shoes"), ("mobile", "Mobile"), ("medical", "Medical"), ("pharma", "Medical"),
]

SHEET_HEADERS = ["Product", "Unit", "Stock", "Selling Price", "Cost Price"]
MAX_UPLOAD_BYTES = 5 * 1024 * 1024
MAX_UPLOAD_ROWS = 2000


def name_key(name):
    """Same rule as the bulk save: case and extra spaces do not make a different product."""
    return " ".join(str(name or "").split()).lower()


def _number(value):
    return float(value) if value is not None else None


def build_sheets(business_type, products):
    """Returns [{name, rows: [{name, unit, stockQty, sellingPrice, costPrice, existing}]}].
    products are the owner's active products as dicts (name, category, unit, stockQty, ...).
    Order: the owner's own categories, then the sheet for their business, then the rest."""
    sheets = {}

    def sheet(name):
        for existing_name in sheets:
            if existing_name.lower() == name.lower():
                return sheets[existing_name]
        sheets[name] = []
        return sheets[name]

    owned = set()
    own_order = []
    for product in products:
        category = (product.get("category") or "General").strip() or "General"
        rows = sheet(category)
        if category.lower() not in [c.lower() for c in own_order]:
            own_order.append(category)
        rows.append({
            "name": product["name"],
            "unit": product.get("unit") or "kg",
            "stockQty": _number(product.get("stockQty")),
            "sellingPrice": _number(product.get("sellingPrice")),
            "costPrice": _number(product.get("costPrice")),
            "existing": True,
        })
        owned.add(name_key(product["name"]))

    for sheet_name, items in DEFAULT_SHEETS.items():
        rows = sheet(sheet_name)
        for name, unit in items:
            # A product the owner already has is shown once, in the sheet of its own category.
            if name_key(name) in owned:
                continue
            rows.append({"name": name, "unit": unit, "stockQty": None, "sellingPrice": None,
                         "costPrice": None, "existing": False})

    bt = (business_type or "").lower()
    business_sheet = next((sheet_name for hint, sheet_name in BUSINESS_SHEET_HINTS if hint in bt), None)
    order = list(own_order)
    if business_sheet and business_sheet not in order:
        order.append(business_sheet)
    order += [name for name in sheets if name not in order]
    return [{"name": name, "rows": sheets[name]} for name in order]


def build_workbook_base64(sheets):
    """The sheets as an .xlsx file (base64), one worksheet per sheet."""
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill
    from openpyxl.worksheet.datavalidation import DataValidation

    workbook = Workbook()
    workbook.remove(workbook.active)
    header_fill = PatternFill("solid", fgColor="E2EFDA")
    for sheet in sheets:
        # Excel limits sheet names to 31 characters and forbids a few symbols.
        title = "".join(ch for ch in sheet["name"] if ch not in '[]:*?/\\')[:31] or "Sheet"
        ws = workbook.create_sheet(title=title)
        ws.append(SHEET_HEADERS)
        for cell in ws[1]:
            cell.font = Font(bold=True)
            cell.fill = header_fill
        for row in sheet["rows"]:
            ws.append([row["name"], row["unit"], row["stockQty"], row["sellingPrice"], row["costPrice"]])
        ws.column_dimensions["A"].width = 28
        for column in "BCDE":
            ws.column_dimensions[column].width = 13
        ws.freeze_panes = "A2"
        positive = DataValidation(type="decimal", operator="greaterThanOrEqual", formula1="0", allow_blank=True)
        ws.add_data_validation(positive)
        positive.add("C2:E5000")
    buffer = io.BytesIO()
    workbook.save(buffer)
    return base64.b64encode(buffer.getvalue()).decode("ascii")


class SheetFileError(ValueError):
    """The uploaded file cannot be read as a product sheet; the message is shown to the owner."""


def _cell_number(value):
    if value in (None, ""):
        return None
    try:
        number = float(str(value).replace(",", "").replace("₹", "").strip())
    except ValueError:
        return None
    return number if number >= 0 else None


def parse_workbook_base64(file_base64):
    """Reads an uploaded .xlsx back into sheets. Columns are found by header name, so the owner
    may reorder them; without recognisable headers the template's column order is assumed."""
    from openpyxl import load_workbook

    try:
        data = base64.b64decode(file_base64, validate=False)
    except (ValueError, TypeError):
        raise SheetFileError("The file could not be read.")
    if len(data) > MAX_UPLOAD_BYTES:
        raise SheetFileError("The file is too large. Please keep it under 5 MB.")
    # Refuse archives that expand to something huge before handing them to the XML parser.
    try:
        with zipfile.ZipFile(io.BytesIO(data)) as archive:
            if sum(info.file_size for info in archive.infolist()) > 50 * 1024 * 1024:
                raise SheetFileError("The file is too large.")
    except zipfile.BadZipFile:
        raise SheetFileError("Please choose an Excel file (.xlsx).")
    try:
        workbook = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    except Exception:
        raise SheetFileError("Please choose an Excel file (.xlsx).")

    aliases = {
        "name": ("product", "product name", "name", "item", "item name"),
        "unit": ("unit", "uom"),
        "stockQty": ("stock", "qty", "quantity", "stock qty", "opening stock"),
        "sellingPrice": ("selling price", "price", "rate", "mrp", "sale price"),
        "costPrice": ("cost price", "cost", "purchase price", "buying price"),
    }
    sheets = []
    total_rows = 0
    for ws in workbook.worksheets:
        rows = []
        columns = None
        for values in ws.iter_rows(values_only=True):
            cells = [str(v).strip() if v is not None else "" for v in values]
            if columns is None:
                lowered = [c.lower() for c in cells]
                found = {field: next((i for i, c in enumerate(lowered) if c in names), None)
                         for field, names in aliases.items()}
                if found["name"] is not None:
                    columns = found
                    continue
                columns = {"name": 0, "unit": 1, "stockQty": 2, "sellingPrice": 3, "costPrice": 4}

            def pick(field):
                index = columns.get(field)
                return values[index] if index is not None and index < len(values) else None

            name = " ".join(str(pick("name") or "").split())
            if not name:
                continue
            rows.append({
                "name": name[:150],
                "unit": str(pick("unit") or "").strip()[:30] or None,
                "stockQty": _cell_number(pick("stockQty")),
                "sellingPrice": _cell_number(pick("sellingPrice")),
                "costPrice": _cell_number(pick("costPrice")),
            })
            total_rows += 1
            if total_rows > MAX_UPLOAD_ROWS:
                raise SheetFileError(f"The file has too many products. Please keep it under {MAX_UPLOAD_ROWS}.")
        if rows:
            sheets.append({"name": ws.title.strip() or "General", "rows": rows})
    workbook.close()
    if not sheets:
        raise SheetFileError("No products were found in the file.")
    return sheets
